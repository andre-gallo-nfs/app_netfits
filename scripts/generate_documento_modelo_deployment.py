import os
import sys
import shutil
import subprocess
from datetime import datetime
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_heading_with_badge(doc, text, badge_text="OK", color_hex="10B981"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    
    run_badge = p.add_run(f"[{badge_text}] ")
    run_badge.font.name = "Arial"
    run_badge.font.bold = True
    run_badge.font.size = Pt(10)
    run_badge.font.color.rgb = RGBColor.from_string(color_hex)
    
    run_text = p.add_run(text)
    run_text.font.name = "Arial Black"
    run_text.font.size = Pt(12)
    run_text.font.color.rgb = RGBColor(15, 23, 42)

def add_callout(doc, text, title="NOTA DE SEGURANÇA", bg_hex="F8FAFC", border_hex="3B82F6"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(7.1)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    # Left border
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    r_title = p.add_run(f"• {title}: ")
    r_title.font.name = "Arial"
    r_title.font.bold = True
    r_title.font.size = Pt(9.5)
    r_title.font.color.rgb = RGBColor.from_string(border_hex)
    
    r_body = p.add_run(text)
    r_body.font.name = "Arial"
    r_body.font.size = Pt(9.5)
    r_body.font.color.rgb = RGBColor(51, 65, 85)

def create_document():
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    # Header Table
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    
    col_left, col_right = header_table.rows[0].cells
    col_left.width = Inches(4.8)
    col_right.width = Inches(2.3)

    p_title = col_left.paragraphs[0]
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_netfits = p_title.add_run("NETFITS PLATAFORMA DIGITAL S.A.\n")
    run_netfits.font.name = "Arial Black"
    run_netfits.font.size = Pt(18)
    run_netfits.font.color.rgb = RGBColor(16, 185, 129)

    run_sub = p_title.add_run("Relatório Técnico de Arquitetura de Software & Implantação")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(10)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    p_meta = col_right.paragraphs[0]
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_meta.paragraph_format.space_before = Pt(0)
    p_meta.paragraph_format.space_after = Pt(2)
    r_meta = p_meta.add_run(
        "DOC REF: NFS-DEP-2026-001\n"
        "DATA: 08/10/2026\n"
        "VERSÃO: 1.0 OFICIAL\n"
        "STATUS: HOMOLOGADO & ATIVO"
    )
    r_meta.font.name = "Courier New"
    r_meta.font.size = Pt(8.5)
    r_meta.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Document Main Title
    p_doc_title = doc.add_paragraph()
    p_doc_title.paragraph_format.space_before = Pt(8)
    p_doc_title.paragraph_format.space_after = Pt(4)
    run_main_title = p_doc_title.add_run("ARQUITETURA DE DEPLOYMENT DUAL (HOMOLOGAÇÃO & PRODUÇÃO) COM PORTÃO DE VALIDAÇÃO HUMANA")
    run_main_title.font.name = "Arial Black"
    run_main_title.font.size = Pt(14)
    run_main_title.font.color.rgb = RGBColor(15, 23, 42)

    # Executive Subtitle
    p_doc_sub = doc.add_paragraph()
    p_doc_sub.paragraph_format.space_before = Pt(0)
    p_doc_sub.paragraph_format.space_after = Pt(10)
    run_main_sub = p_doc_sub.add_run(
        "Dossiê técnico e executivo detalhando o diagnóstico da solicitação, a fundamentação da arquitetura proposta, "
        "a implementação completa dos ambientes segregados e o protocolo definitivo de governança e tombamento para produção."
    )
    run_main_sub.font.name = "Arial"
    run_main_sub.font.size = Pt(10)
    run_main_sub.font.italic = True
    run_main_sub.font.color.rgb = RGBColor(71, 85, 105)

    add_callout(
        doc,
        "A partir de 08/10/2026, nenhuma alteração, refatoração ou nova funcionalidade desenvolvida por inteligência artificial "
        "ou desenvolvedores é publicada diretamente no ambiente de Produção. Todo ciclo opera exclusivamente em Homologação, "
        "condicionado à validação soberana do CEO/Owner antes de qualquer tombamento.",
        title="DIRETRIZ DE BLINDAGEM OPERACIONAL MANDATÓRIA",
        bg_hex="FEF3C7",
        border_hex="D97706"
    )

    # SEÇÃO 1
    add_heading_with_badge(doc, "1. DIAGNÓSTICO E CONTEXTO DA SOLICITAÇÃO", "DIAGNÓSTICO", "2563EB")
    
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.add_run(
        "A plataforma Netfits atingiu recentemente marcos decisivos de maturidade comercial e técnica, incluindo a homologação oficial "
        "das chaves criptográficas RSA RS256 da Rock Encantech (StoreId: RhOFkbZJIN), a validação da cotação de fidelidade (R$ 0,01 por ponto NFS) "
        "e o recebimento de pedidos reais com apuração de 6% de Take-Rate. Neste estágio de operação real com parceiros comerciais e atletas, "
        "o modelo de implantação contínua anterior apresentava um risco estrutural inaceitável."
    )
    
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.add_run(
        "O modelo anterior operava sob push direto na branch principal ('main'), onde qualquer ajuste disparava compilação e publicação "
        "imediata em produção. Embora dotado de 'robôs de qualidade' (verificações de rotas, linters e checagem de tipos TypeScript via tsc --noEmit), "
        "identificou-se que esses mecanismos possuem um ponto cego crítico:\n\n"
        "1. Os robôs de teste automatizados apenas validam a sintaxe e a integridade de compilação do código. Se o código compilar sem erros de digitação, os testes são aprovados com 100% de sucesso.\n"
        "2. Os robôs não avaliam a interpretação da intenção de negócio: caso a IA compreenda um requisito de forma equivocada no prompt, ou execute uma solução sintaticamente perfeita porém funcionalmente divergente (ex: apagar dados legítimos de perfil de cliente ou alterar inadvertidamente cálculos de regras de pontuação), o código defeituoso seria promovido direto para os usuários finais e para a Rock Encantech.\n"
        "3. Foi exatamente essa a causa raiz de incidentes recentes observados em sessões anteriores (como a sobrescrita indevida de dados de perfil)."
    )

    # SEÇÃO 2
    add_heading_with_badge(doc, "2. A SOLUÇÃO ARQUITETURAL: AMBIENTE DUAL & PORTÃO HUMANO", "ARQUITETURA", "10B981")
    
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.add_run(
        "Para erradicar qualquer possibilidade de erro funcional em produção, foi desenhada e implementada a arquitetura "
        "de Ambientes Segregados com Portão Humano (Human-in-the-Loop Gate). O modelo desacopla completamente o espaço de "
        "experimentação, validação e refinamento do ambiente oficial de clientes."
    )

    # Tabela Comparativa de Ambientes
    tbl_amb = doc.add_table(rows=7, cols=3)
    tbl_amb.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_amb.autofit = False

    headers = ["Dimensão Operacional", "🧪 Homologação (Staging)", "🚀 Produção Oficial (Prod)"]
    widths = [Inches(2.2), Inches(2.45), Inches(2.45)]
    
    for i, h in enumerate(headers):
        cell = tbl_amb.rows[0].cells[i]
        cell.width = widths[i]
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    data_rows = [
        ("Branch do Repositório Git", "staging (ramo de trabalho)", "main (ramo protegido bloqueado)"),
        ("Identificação Visual no App", "Badge discreta 'HOMOLOG' no topo", "Sem badge (visual limpo e definitivo)"),
        ("Arquivo de Configuração", ".env.staging (VITE_APP_ENV=staging)", ".env.production (VITE_APP_ENV=production)"),
        ("Integração Rock Encantech", "Ambiente Sandbox / Simulação", "Chaves Oficiais Produção (PUQ4cwt2...)"),
        ("Base de Dados & Carteiras", "Bancos/stores de teste isolados", "Ledger oficial criptografado dos clientes"),
        ("Autorização de Publicação", "Deploy automático a cada demanda", "Somente pelo André após teste humano")
    ]

    for idx, (dim, stg, prd) in enumerate(data_rows, start=1):
        row = tbl_amb.rows[idx]
        bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate([dim, stg, prd]):
            cell = row.cells[c_idx]
            cell.width = widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            if c_idx == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(30, 41, 59)
            elif c_idx == 1:
                r.font.color.rgb = RGBColor(180, 83, 9)
            else:
                r.font.color.rgb = RGBColor(21, 128, 61)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # SEÇÃO 3
    add_heading_with_badge(doc, "3. DETALHAMENTO DA IMPLEMENTAÇÃO TÉCNICA", "IMPLEMENTAÇÃO", "8B5CF6")
    
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.add_run(
        "A implantação foi executada diretamente na raiz da arquitetura do repositório 'app_netfits', cobrindo 6 pilares de engenharia:"
    )

    items_impl = [
        ("3.1 Criação e Vinculação da Branch 'staging':", 
         "Criada a ramificação oficial 'staging' sincronizada com o GitHub (origin/staging). Todo desenvolvimento ativo passa a residir nela. A branch 'main' foi mantida em seu estado estável (commit 77d942d)."),
        ("3.2 Segregação de Ambientes (.env.staging e .env.production):", 
         "Criados arquivos formais de variáveis com injeção automática em tempo de build. Homologação roda com flags de sandbox e GA de teste; Produção roda com parâmetros oficiais e segredos protegidos na nuvem."),
        ("3.3 Feedback Visual Dinâmico no Cabeçalho (AppShell):", 
         "Atualizado o componente 'src/components/AppShell.tsx' para detectar 'VITE_APP_ENV === staging'. Em homologação, o app exibe uma etiqueta âmbar 'HOMOLOG' ao lado do logotipo Netfits. Em produção, a tag não é renderizada."),
        ("3.4 Quality Gate Automatizado no GitHub Actions (ci-gate.yml):", 
         "Implementado workflow '.github/workflows/ci-gate.yml'. Ele dispara validação em push para staging, main e pull requests, executando validação de rotas, tsc e build de produção."),
        ("3.5 Script de Tombamento com 9 Travas de Segurança (promote:prod):", 
         "Desenvolvido o utilitário 'scripts/promote-to-production.mjs' (acessível via 'npm run promote:prod'). O script impede execuções forasteiras, valida pendências, testa o build, realiza merge limpo sem fast-forward, publica na main e devolve o desenvolvedor em segurança para a staging."),
        ("3.6 Blindagem Institucional no GEMINI.md e Guia Operacional:", 
         "Revogada a diretriz antiga de deploy contínuo em main. Registrado no 'GEMINI.md' a proibição absoluta de push direto em produção por qualquer IA ou desenvolvedor, assegurando conformidade perene.")
    ]

    for title, desc in items_impl:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(4)
        r_t = p.add_run(f"{title} ")
        r_t.font.name = "Arial"
        r_t.font.bold = True
        r_t.font.size = Pt(9.5)
        r_t.font.color.rgb = RGBColor(30, 41, 59)
        
        r_d = p.add_run(desc)
        r_d.font.name = "Arial"
        r_d.font.size = Pt(9.5)
        r_d.font.color.rgb = RGBColor(71, 85, 105)

    # SEÇÃO 4
    add_heading_with_badge(doc, "4. O FLUXO DE OPERAÇÃO DIÁRIA (PASSO A PASSO)", "FLUXO", "0EA5E9")
    
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.add_run(
        "A partir desta implementação, o ciclo de qualquer nova solicitação do André seguirá o seguinte roteiro estrito:\n\n"
        "• Passo 1 (Demanda): O André envia o pedido ou melhoria desejada via prompt.\n"
        "• Passo 2 (Desenvolvimento em Staging): A IA implementa e valida os códigos exclusivamente na branch 'staging' (push para origin staging).\n"
        "• Passo 3 (Disponibilização em Homologação): O sistema sobe automaticamente na URL de homologação (com a tag visual 'HOMOLOG').\n"
        "• Passo 4 (Auditoria do André): O André acessa o link e valida se a IA entendeu o prompt e se a execução não causou desvios.\n"
        "• Passo 5 (Correção em caso de reprovação): Caso identifique qualquer incongruência, a correção é feita em Homologação com a Produção 100% segura e intocada.\n"
        "• Passo 6 (Autorização & Tombamento Oficial): Uma vez aprovado pelo André com a frase 'Pode tombar para produção', o comando 'npm run promote:prod' é disparado, promovendo a versão para 'main' e atualizando a produção oficial."
    )

    # SEÇÃO 5
    add_heading_with_badge(doc, "5. CONCLUSÃO E CERTIFICAÇÃO DE GOVERNANÇA", "CERTIFICAÇÃO", "10B981")
    
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.add_run(
        "Com esta entrega, a Netfits Plataforma Digital S.A. elimina o risco de quebras não planejadas em produção causadas por desvios de interpretação "
        "de prompts ou alterações intempestivas de código. A plataforma atinge o padrão de excelência de governança de software exigido por investidores, "
        "garantindo estabilidade contínua para os parceiros corporativos (Rock Encantech) e atletas."
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Assinatura
    p_sign = doc.add_paragraph()
    p_sign.paragraph_format.space_before = Pt(14)
    p_sign.paragraph_format.space_after = Pt(0)
    r_s = p_sign.add_run(
        "_____________________________________________________________________\n"
        "André Gallo | Diretor Executivo & Fundador (CEO / Owner)\n"
        "Antigravity AI | Engenharia de Arquitetura de Software & FinOps\n"
        "Netfits Plataforma Digital S.A. — Certificação de Produção 2026"
    )
    r_s.font.name = "Arial"
    r_s.font.size = Pt(8.5)
    r_s.font.color.rgb = RGBColor(100, 116, 139)
    r_s.font.italic = True

    return doc

def main():
    print("[1/3] Gerando documento Word oficial do Modelo de Deployment...")
    doc = create_document()
    
    output_dir_local = r"C:\Users\aacga\Projetos\app_netfits"
    output_dir_onedrive = r"c:\Users\aacga\OneDrive\netfits"
    output_dir_artifacts = r"C:\Users\aacga\.gemini\antigravity\brain\0c50996a-827d-4c1d-b276-5ffdfd4f7f85"

    docx_filename = "NETFITS_Dossie_Oficial_Modelo_Deployment_Homologacao_e_Producao_08_10_2026.docx"
    pdf_filename = "NETFITS_Dossie_Oficial_Modelo_Deployment_Homologacao_e_Producao_08_10_2026.pdf"

    docx_path = os.path.join(output_dir_local, docx_filename)
    pdf_path = os.path.join(output_dir_local, pdf_filename)

    doc.save(docx_path)
    print(f"Docx salvo com sucesso: {docx_path}")

    print("[2/3] Convertendo Docx para PDF via Word COM Automation...")
    ps_cmd = f"""
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    $doc = $word.Documents.Open('{docx_path}')
    $doc.SaveAs([ref]'{pdf_path}', [ref]17)
    $doc.Close()
    $word.Quit()
    """
    res = subprocess.run(["powershell", "-Command", ps_cmd], capture_output=True, text=True)
    if not os.path.exists(pdf_path):
        print(f"Erro na conversão PDF: {res.stderr}")
        sys.exit(1)
    
    print(f"PDF gerado com sucesso: {pdf_path} ({os.path.getsize(pdf_path)} bytes)")

    print("[3/3] Replicando cópias para OneDrive e Artefatos...")
    shutil.copy2(pdf_path, os.path.join(output_dir_onedrive, pdf_filename))
    shutil.copy2(docx_path, os.path.join(output_dir_onedrive, docx_filename))
    shutil.copy2(pdf_path, os.path.join(output_dir_artifacts, pdf_filename))
    shutil.copy2(docx_path, os.path.join(output_dir_artifacts, docx_filename))
    print("Concluído com 100% de sucesso!")

if __name__ == "__main__":
    main()
