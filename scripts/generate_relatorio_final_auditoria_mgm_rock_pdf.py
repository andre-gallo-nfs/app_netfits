import os
import sys
import shutil
import subprocess
from datetime import datetime
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_heading_with_badge(doc, text, badge_text="OK", color_hex="10B981"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    
    run_badge = p.add_run(f"[{badge_text}] ")
    run_badge.font.name = "Arial"
    run_badge.font.bold = True
    run_badge.font.size = Pt(9.5)
    run_badge.font.color.rgb = RGBColor.from_string(color_hex)
    
    run_text = p.add_run(text)
    run_text.font.name = "Arial Black"
    run_text.font.size = Pt(12)
    run_text.font.color.rgb = RGBColor(15, 23, 42)

def create_document():
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    
    col_left, col_right = header_table.rows[0].cells
    col_left.width = Inches(4.8)
    col_right.width = Inches(2.3)

    p_title = col_left.paragraphs[0]
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_netfits = p_title.add_run("NETFITS\n")
    run_netfits.font.name = "Arial Black"
    run_netfits.font.size = Pt(20)
    run_netfits.font.color.rgb = RGBColor(16, 185, 129)

    run_sub = p_title.add_run("RELATÓRIO OFICIAL DE AUDITORIA & ADEQUAÇÃO DE REGRAS MGM\n")
    run_sub.font.name = "Arial"
    run_sub.font.bold = True
    run_sub.font.size = Pt(10)
    run_sub.font.color.rgb = RGBColor(15, 23, 42)

    run_desc = p_title.add_run("Ajuste de Comissões de Compra, Homologação de Webhooks Rock & Estorno Contábil")
    run_desc.font.name = "Arial"
    run_desc.font.size = Pt(8.5)
    run_desc.font.color.rgb = RGBColor(100, 116, 139)

    p_meta = col_right.paragraphs[0]
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_meta.paragraph_format.space_before = Pt(0)
    p_meta.paragraph_format.space_after = Pt(0)
    run_meta = p_meta.add_run(
        "DATA: 07/10/2026\n"
        "VERSÃO: 3.2 PRODUÇÃO\n"
        "STATUS: HOMOLOGADO\n"
        "AMBIENTE: LIVE PROD"
    )
    run_meta.font.name = "Arial"
    run_meta.font.size = Pt(8)
    run_meta.font.bold = True
    run_meta.font.color.rgb = RGBColor(71, 85, 105)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(6)
    p_div.paragraph_format.space_after = Pt(10)
    r_div = p_div.add_run("—" * 65)
    r_div.font.color.rgb = RGBColor(226, 232, 240)

    # 1. DIRETRIZ ESTRATÉGICA MGM
    add_heading_with_badge(doc, "1. DIRETRIZ DE NEGÓCIO: PROGRAMA MGM (MEMBER GET MEMBER)", "DIRETRIZ OFICIAL", "2563EB")
    p_mgm = doc.add_paragraph()
    p_mgm.paragraph_format.space_after = Pt(6)
    r_mgm = p_mgm.add_run(
        "Conforme deliberação da liderança executiva da Netfits, foram formalizadas as seguintes diretrizes para o programa de indicação (MGM):\n\n"
        "• Indicação Efetivada no Cadastro (Ativa): O usuário indicador recebe a bonificação fixa por novo usuário que conclui seu cadastro utilizando seu código de indicação (ex: Carlos recebeu 50 nfs pela indicação efetivada de Cristiane).\n"
        "• Comissão de Compras no Marketplace (Desativada no Momento): Não há vigência de comissão percentual (5%) sobre compras de indicados no momento. Esta modalidade entrará em vigor futuramente e será restrita exclusivamente aos usuários assinantes do Clube Netfits.\n"
        "• Ajuste Contábil Realizado: A comissão provisória de 25 nfs referente à compra de Cristiane (#SOP0711045469) foi totalmente revogada da conta de Carlos Rodrigo Formigari, retornando seu saldo legítimo para exatamente 110 nfs."
    )
    r_mgm.font.name = "Arial"
    r_mgm.font.size = Pt(9.5)
    r_mgm.font.color.rgb = RGBColor(51, 65, 85)

    # 2. STATUS DAS 3 CORREÇÕES SOLICITADAS PELA ROCK
    add_heading_with_badge(doc, "2. STATUS DAS CORREÇÕES TÉCNICAS (FEEDBACK EDINHO / ROCK)", "100% CUMPRIDO", "10B981")
    
    t_rock = doc.add_table(rows=4, cols=3)
    t_rock.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_rock.autofit = False

    headers = ["Exigência Rock", "Ação Netfits Implementada", "Validação em Produção"]
    widths = [Inches(1.8), Inches(3.3), Inches(1.9)]

    for idx, text in enumerate(headers):
        cell = t_rock.cell(0, idx)
        cell.width = widths[idx]
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    itens_rock = [
        (
            "1. Proteção LGPD em GET /api/orders",
            "Exigência estrita de x-api-key no método GET. Chamadas desautenticadas retornam HTTP 401 Unauthorized sem dados de pedidos. Chamadas autenticadas possuem CPF e e-mail totalmente omitidos da serialização.",
            "curl sem chave: HTTP 401\ncurl com chave: HTTP 200 (sem dados sensíveis)",
            "F8FAFC"
        ),
        (
            "2. Estorno Contábil & Limpeza do Banco",
            "Estorno dos 780 nfs indevidos de Carlos (PFM) e 1.080 nfs superestimados de André (GTJ). Registros PFM e GTJ removidos do banco persistente, mantendo unicamente SOP0711045469 para receber o payload oficial da Rock.",
            "Gist GitHub sincronizado.\nBase pronta para replay como autoridade máxima.",
            "FFFFFF"
        ),
        (
            "3. Idempotência por _id & Upsert de Status",
            "Travas idempotentes a nível de banco e ledger: cada _id gera no máximo 1 crédito e 1 débito. Upsert hierárquico preserva cronologia de updatedAt e prioridade de status.",
            "Testes com disparos repetidos confirmam isReplay: true sem duplicidade.",
            "F8FAFC"
        )
    ]

    for row_idx, item in enumerate(itens_rock, start=1):
        for col_idx in range(3):
            cell = t_rock.cell(row_idx, col_idx)
            cell.width = widths[col_idx]
            set_cell_background(cell, item[3])
            set_cell_margins(cell, 80, 80, 100, 100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(item[col_idx])
            r.font.name = "Arial"
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(30, 41, 59)

    # 3. SALDOS AUDITADOS DA LIDERANÇA
    add_heading_with_badge(doc, "3. AUDITORIA CONTÁBIL DOS SALDOS EM PRODUÇÃO", "AUDITADO", "10B981")
    
    t_bal = doc.add_table(rows=4, cols=4)
    t_bal.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_bal.autofit = False

    bal_headers = ["Usuário / ID", "Saldo Atual (Pré-Replay)", "Composição do Saldo Atual", "Saldo Projetado Pós-Replay"]
    bal_widths = [Inches(1.8), Inches(1.4), Inches(2.3), Inches(1.5)]

    for idx, text in enumerate(bal_headers):
        cell = t_bal.cell(0, idx)
        cell.width = bal_widths[idx]
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    saldos_data = [
        (
            "André Gallo\n(usr_andre)",
            "50 nfs",
            "50 nfs (Bônus de Boas-Vindas)",
            "1.091 nfs\n(+227 PFM e +814 GTJ após replay da Rock)",
            "FFFFFF"
        ),
        (
            "Carlos Formigari\n(usr_carlos_formigari)",
            "110 nfs",
            "50 boas-vindas + 50 indicação Cristiane + 10 feed (sem compras e sem comissão MGM)",
            "110 nfs\n(Intacto — comissão de compras MGM desativada)",
            "F8FAFC"
        ),
        (
            "Cristiane Formigari\n(user-1791370530242)",
            "564 nfs",
            "50 boas-vindas + 50 feed - 50 uso shop + 514 cashback compra SOP",
            "564 nfs\n(Intacto — pedido já auditado e liquidado)",
            "FFFFFF"
        )
    ]

    for row_idx, item in enumerate(saldos_data, start=1):
        for col_idx in range(4):
            cell = t_bal.cell(row_idx, col_idx)
            cell.width = bal_widths[col_idx]
            set_cell_background(cell, item[4])
            set_cell_margins(cell, 80, 80, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(item[col_idx])
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(15, 23, 42)

    # 4. TERMO DE HOMOLOGAÇÃO
    add_heading_with_badge(doc, "4. CERTIFICAÇÃO DE PRONTIDÃO PARA REPLAY DA ROCK", "HOMOLOGADO", "10B981")
    p_cert = doc.add_paragraph()
    r_cert = p_cert.add_run(
        "Certificamos que o ambiente de produção da Netfits Tecnologia Ltda. está 100% pronto e auditado para o recebimento do payload oficial dos pedidos GTJ0522372096, PFM0610443019 e SOP0711045469 via POST https://www.netfits.com.br/api/orders.\n\n"
        "O replay do time da Rock Encantech passará a ser a autoridade absoluta no armazenamento persistente em nuvem da Netfits."
    )
    r_cert.font.name = "Arial"
    r_cert.font.size = Pt(9)
    r_cert.font.color.rgb = RGBColor(71, 85, 105)

    p_sign = doc.add_paragraph()
    p_sign.paragraph_format.space_before = Pt(16)
    p_sign.paragraph_format.space_after = Pt(0)
    r_s = p_sign.add_run(
        "_____________________________________________________________________\n"
        "Netfits Tecnologia Ltda. | Diretoria de Engenharia & Operações\n"
        "Google Cloud Startup Program — Certificação de Produção 2026"
    )
    r_s.font.name = "Arial"
    r_s.font.size = Pt(7.5)
    r_s.font.color.rgb = RGBColor(148, 163, 184)
    r_s.font.italic = True

    return doc

def main():
    print("[1/3] Gerando documento Word com auditoria MGM e Webhooks...")
    doc = create_document()
    
    output_dir_local = r"C:\Users\aacga\Projetos\app_netfits"
    output_dir_onedrive = r"c:\Users\aacga\OneDrive\netfits"
    output_dir_artifacts = r"C:\Users\aacga\.gemini\antigravity\brain\5ce19ed4-336b-40b0-99dc-9784486f1a69"

    docx_filename = "NETFITS_Relatorio_Final_Auditoria_MGM_e_Webhooks_Rock_07_10_2026.docx"
    pdf_filename = "NETFITS_Relatorio_Final_Auditoria_MGM_e_Webhooks_Rock_07_10_2026.pdf"

    docx_path = os.path.join(output_dir_local, docx_filename)
    pdf_path = os.path.join(output_dir_local, pdf_filename)

    doc.save(docx_path)
    print(f"Docx salvo com sucesso: {docx_path}")

    print("[2/3] Convertendo Docx para PDF via Word COM Automation...")
    ps_cmd = f"""
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    $doc = $word.Documents.Open('{docx_path}')
    $doc.SaveAs([ref]'{pdf_path}', [ref]17)
    $doc.Close()
    $word.Quit()
    """
    res = subprocess.run(["powershell", "-Command", ps_cmd], capture_output=True, text=True)
    if not os.path.exists(pdf_path):
        print(f"Erro na conversão PDF: {res.stderr}")
        sys.exit(1)
    
    print(f"PDF gerado com sucesso: {pdf_path} ({os.path.getsize(pdf_path)} bytes)")

    print("[3/3] Replicando cópias para OneDrive e Artefatos...")
    shutil.copy2(pdf_path, os.path.join(output_dir_onedrive, pdf_filename))
    shutil.copy2(docx_path, os.path.join(output_dir_onedrive, docx_filename))
    shutil.copy2(pdf_path, os.path.join(output_dir_artifacts, pdf_filename))
    print("Concluído com 100% de sucesso!")

if __name__ == "__main__":
    main()
