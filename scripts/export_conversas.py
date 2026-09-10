import os
import json
import re
import shutil
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def clean_user_text(raw):
    text = raw.strip()
    # Remove USER_REQUEST tags
    m = re.search(r"<USER_REQUEST>(.*?)</USER_REQUEST>", text, re.DOTALL)
    if m:
        text = m.group(1).strip()
    # Ignore pure internal messages
    if text.startswith("<CONTEXT_SUMMARY>") or text.startswith("<SYSTEM_MESSAGE>"):
        return None
    # Remove trailing metadata
    text = re.sub(r"<ADDITIONAL_METADATA>.*?</ADDITIONAL_METADATA>", "", text, flags=re.DOTALL).strip()
    return text if text else None

def clean_assistant_text(raw):
    text = raw.strip()
    if not text:
        return None
    if text.startswith("Created At:") or "jsonhook__" in text:
        return None
    # Remove system notice tags
    text = re.sub(r"<SYSTEM_MESSAGE>.*?</SYSTEM_MESSAGE>", "", text, flags=re.DOTALL).strip()
    return text if text else None

def parse_transcript(log_path):
    dialogues = []
    current_user = None
    current_responses = []

    with open(log_path, "r", encoding="utf-8") as f:
        for line in f:
            try:
                item = json.loads(line)
            except Exception:
                continue

            stype = item.get("type")
            content = item.get("content", "")
            created_at = item.get("created_at", "")

            if stype == "USER_INPUT":
                cleaned_u = clean_user_text(content)
                if cleaned_u:
                    if current_user and current_responses:
                        dialogues.append({
                            "user": current_user["text"],
                            "time": current_user["time"],
                            "assistant": "\n\n".join(current_responses)
                        })
                    current_user = {"text": cleaned_u, "time": created_at}
                    current_responses = []
            elif stype == "PLANNER_RESPONSE":
                cleaned_a = clean_assistant_text(content)
                if cleaned_a:
                    current_responses.append(cleaned_a)

    if current_user and current_responses:
        dialogues.append({
            "user": current_user["text"],
            "time": current_user["time"],
            "assistant": "\n\n".join(current_responses)
        })

    return dialogues

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def generate_docx(dialogues, out_path):
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    NAVY = RGBColor(15, 23, 42)
    PURPLE = RGBColor(124, 58, 237)
    GREEN = RGBColor(16, 185, 129)
    DARK = RGBColor(30, 41, 59)
    GRAY = RGBColor(100, 116, 139)

    p_top = doc.add_paragraph()
    r1 = p_top.add_run("NETFITS — REGISTRO HISTÓRICO INTEGRAL DE CONVERSAS & DIRETRIZES\n")
    r1.font.name = "Arial"
    r1.font.size = Pt(11)
    r1.font.bold = True
    r1.font.color.rgb = PURPLE

    r2 = p_top.add_run("Histórico Oficial das Interações, Prompts, Decisões de Arquitetura e Engenharia")
    r2.font.name = "Arial"
    r2.font.size = Pt(20)
    r2.font.bold = True
    r2.font.color.rgb = NAVY

    p_sub = doc.add_paragraph()
    rs = p_sub.add_run(f"Exportação completa realizada em {datetime.now().strftime('%d/%m/%Y às %H:%M:%S')} | Total de Interações: {len(dialogues)}")
    rs.font.name = "Arial"
    rs.font.size = Pt(10)
    rs.font.italic = True
    rs.font.color.rgb = GRAY

    doc.add_paragraph()

    for idx, turn in enumerate(dialogues, 1):
        # User Box
        tbl_u = doc.add_table(rows=1, cols=1)
        tbl_u.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_u.autofit = False
        cell_u = tbl_u.rows[0].cells[0]
        cell_u.width = Inches(6.5)
        set_cell_background(cell_u, "F8FAFC")
        set_cell_margins(cell_u, 120, 120, 160, 160)
        pu = cell_u.paragraphs[0]
        ru_tag = pu.add_run(f"👤 INTERAÇÃO #{idx} — ANDRÉ GALLO (USUÁRIO) [{turn.get('time', '')}]\n")
        ru_tag.font.name = "Arial"
        ru_tag.font.size = Pt(9.5)
        ru_tag.font.bold = True
        ru_tag.font.color.rgb = PURPLE

        ru_txt = pu.add_run(turn["user"])
        ru_txt.font.name = "Arial"
        ru_txt.font.size = Pt(10)
        ru_txt.font.bold = True
        ru_txt.font.color.rgb = NAVY

        # Assistant Box
        doc.add_paragraph()
        tbl_a = doc.add_table(rows=1, cols=1)
        tbl_a.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_a.autofit = False
        cell_a = tbl_a.rows[0].cells[0]
        cell_a.width = Inches(6.5)
        set_cell_background(cell_a, "FFFFFF")
        set_cell_margins(cell_a, 120, 120, 160, 160)
        pa = cell_a.paragraphs[0]
        ra_tag = pa.add_run("🤖 RESPOSTA & DIRETRIZES — ANTIGRAVITY (ASSISTENTE IA)\n\n")
        ra_tag.font.name = "Arial"
        ra_tag.font.size = Pt(9.5)
        ra_tag.font.bold = True
        ra_tag.font.color.rgb = GREEN

        # Keep response reasonably sized in Word if huge
        resp = turn["assistant"]
        if len(resp) > 3000:
            resp = resp[:3000] + "\n\n[... Resposta técnica detalhada com código preservada integralmente no arquivo Markdown ...]"

        ra_txt = pa.add_run(resp)
        ra_txt.font.name = "Arial"
        ra_txt.font.size = Pt(9.5)
        ra_txt.font.color.rgb = DARK

        doc.add_paragraph()

    doc.save(out_path)
    print(f"Word gerado com sucesso: {out_path}")

def main():
    log_path = r"C:\Users\aacga\.gemini\antigravity\brain\9b14c41d-3560-4e77-a0bc-e83c84aad2d0\.system_generated\logs\transcript_full.jsonl"
    if not os.path.exists(log_path):
        log_path = r"C:\Users\aacga\.gemini\antigravity\brain\9b14c41d-3560-4e77-a0bc-e83c84aad2d0\.system_generated\logs\transcript.jsonl"

    project_dir = r"C:\Users\aacga\Projetos\app_netfits"
    onedrive_dir = r"C:\Users\aacga\OneDrive\netfits"

    dest_project = os.path.join(project_dir, "conversas")
    dest_onedrive = os.path.join(onedrive_dir, "conversas")
    dest_raw_project = os.path.join(dest_project, "dados_brutos")
    dest_raw_onedrive = os.path.join(dest_onedrive, "dados_brutos")

    for d in [dest_project, dest_onedrive, dest_raw_project, dest_raw_onedrive]:
        os.makedirs(d, exist_ok=True)

    print("Parseando histórico completo das conversas...")
    dialogues = parse_transcript(log_path)
    print(f"Total de turnos de diálogo recuperados: {len(dialogues)}")

    # 1. Gerar Markdown
    md_content = ["# NETFITS — HISTÓRICO INTEGRAL DE CONVERSAS & DIRETRIZES TÉCNICAS\n"]
    md_content.append(f"> **Data de Exportação:** {datetime.now().strftime('%d/%m/%Y às %H:%M:%S')}")
    md_content.append(f"> **Repositório:** `https://github.com/andre-gallo-nfs/app_netfits.git`")
    md_content.append(f"> **Total de Interações:** {len(dialogues)}\n")
    md_content.append("---\n")

    for idx, turn in enumerate(dialogues, 1):
        md_content.append(f"## Interação #{idx} — [{turn.get('time', '')}]\n")
        md_content.append(f"### 👤 André Gallo (Usuário):\n")
        md_content.append(f"```text\n{turn['user']}\n```\n")
        md_content.append(f"### 🤖 Antigravity (Assistente):\n")
        md_content.append(f"{turn['assistant']}\n")
        md_content.append("\n---\n")

    full_md = "\n".join(md_content)

    md_project = os.path.join(dest_project, "Historico_Completo_Conversas_Netfits.md")
    md_onedrive = os.path.join(dest_onedrive, "Historico_Completo_Conversas_Netfits.md")

    with open(md_project, "w", encoding="utf-8") as f:
        f.write(full_md)
    with open(md_onedrive, "w", encoding="utf-8") as f:
        f.write(full_md)
    print(f"Markdown salvo em: {md_project}")

    # 2. Gerar Word
    docx_project = os.path.join(dest_project, "Historico_Completo_Conversas_Netfits.docx")
    docx_onedrive = os.path.join(dest_onedrive, "Historico_Completo_Conversas_Netfits.docx")
    generate_docx(dialogues, docx_project)
    shutil.copyfile(docx_project, docx_onedrive)
    print(f"Word copiado para OneDrive: {docx_onedrive}")

    # 3. Copiar Dados Brutos (JSONL)
    src_logs = os.path.dirname(log_path)
    for fname in ["transcript.jsonl", "transcript_full.jsonl"]:
        fsrc = os.path.join(src_logs, fname)
        if os.path.exists(fsrc):
            shutil.copyfile(fsrc, os.path.join(dest_raw_project, fname))
            shutil.copyfile(fsrc, os.path.join(dest_raw_onedrive, fname))
            print(f"Log bruto {fname} copiado para dados_brutos.")

    print("Importação concluída com 100% de sucesso!")

if __name__ == "__main__":
    main()
